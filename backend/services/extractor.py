from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List
import re


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


def clean_text(text: str) -> str:
    """Normalize extracted text while preserving useful line breaks."""
    if not text:
        return ""

    text = text.replace("\x00", " ")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def guess_section(text: str) -> str:
    """Best-effort heading detection for evidence metadata."""
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue

        if len(line) <= 100 and (
            line.isupper()
            or re.match(r"^(section|chapter|part|unit|article)\b", line, re.I)
            or re.match(r"^\d+(\.\d+)*[.)]?\s+\S+", line)
        ):
            return line

    return ""


def extract_pdf(path: str, filename: str) -> List[Dict[str, Any]]:
    records: List[Dict[str, Any]] = []

    try:
        import pymupdf
    except ImportError as exc:
        raise RuntimeError(
            "PyMuPDF is missing. Run: python -m pip install pymupdf"
        ) from exc

    try:
        document = pymupdf.open(path)
    except Exception as exc:
        raise RuntimeError(f"Could not open PDF: {exc}") from exc

    try:
        for page_number, page in enumerate(document, start=1):
            try:
                text = page.get_text("text") or ""
            except Exception:
                text = ""

            text = clean_text(text)
            if text:
                records.append(
                    {
                        "document": filename,
                        "page": page_number,
                        "section": guess_section(text) or f"Page {page_number}",
                        "text": text,
                    }
                )
    finally:
        document.close()

    return records


def extract_docx(path: str, filename: str) -> List[Dict[str, Any]]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError(
            "python-docx is missing. Run: python -m pip install python-docx"
        ) from exc

    document = Document(path)
    records: List[Dict[str, Any]] = []
    current_section = "Document"
    paragraph_buffer: List[str] = []

    def flush_paragraphs() -> None:
        nonlocal paragraph_buffer
        if not paragraph_buffer:
            return

        text = "\n".join(paragraph_buffer).strip()
        if text:
            records.append(
                {
                    "document": filename,
                    "page": None,
                    "section": current_section,
                    "text": text,
                }
            )
        paragraph_buffer = []

    for paragraph in document.paragraphs:
        text = clean_text(paragraph.text)
        if not text:
            continue

        style_name = (paragraph.style.name or "").lower()
        if "heading" in style_name or style_name.startswith("title"):
            flush_paragraphs()
            current_section = text
        else:
            paragraph_buffer.append(text)

    flush_paragraphs()

    # Capture tables because rules/circulars often contain important evidence in tables.
    for table_index, table in enumerate(document.tables, start=1):
        rows: List[str] = []
        for row in table.rows:
            cells = [clean_text(cell.text) for cell in row.cells]
            cells = [cell for cell in cells if cell]
            if cells:
                rows.append(" | ".join(cells))

        if rows:
            records.append(
                {
                    "document": filename,
                    "page": None,
                    "section": f"Table {table_index}",
                    "text": "\n".join(rows),
                }
            )

    return records


def extract_txt(path: str, filename: str) -> List[Dict[str, Any]]:
    encodings = ("utf-8", "utf-8-sig", "cp1252", "latin-1")
    text = None

    for encoding in encodings:
        try:
            text = Path(path).read_text(encoding=encoding)
            break
        except UnicodeDecodeError:
            continue

    if text is None:
        raise RuntimeError("Could not decode the TXT file.")

    text = clean_text(text)
    if not text:
        return []

    return [
        {
            "document": filename,
            "page": None,
            "section": guess_section(text) or "Document",
            "text": text,
        }
    ]


def extract_pages(path: str, filename: str) -> List[Dict[str, Any]]:
    """Extract readable content from PDF, DOCX or TXT."""
    file_path = Path(path)
    suffix = file_path.suffix.lower()

    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    if suffix == ".pdf":
        return extract_pdf(str(file_path), filename)
    if suffix == ".docx":
        return extract_docx(str(file_path), filename)
    if suffix == ".txt":
        return extract_txt(str(file_path), filename)

    raise ValueError("Unsupported format. Use PDF, DOCX, or TXT.")


def chunk_records(
    records: List[Dict[str, Any]],
    chunk_size: int = 900,
    overlap: int = 150,
) -> List[Dict[str, Any]]:
    """Split extracted records into overlapping searchable chunks."""
    if chunk_size <= 0 or overlap < 0 or overlap >= chunk_size:
        raise ValueError("Invalid chunk_size/overlap configuration.")

    chunks: List[Dict[str, Any]] = []

    for record in records:
        text = clean_text(str(record.get("text", "")))
        if not text:
            continue

        start = 0
        while start < len(text):
            end = min(start + chunk_size, len(text))

            # Prefer a paragraph/line boundary when possible.
            if end < len(text):
                boundary = text.rfind("\n", start, end)
                if boundary >= start + int(chunk_size * 0.55):
                    end = boundary

            chunk_text = text[start:end].strip()
            if chunk_text:
                chunks.append(
                    {
                        "document": record.get("document", ""),
                        "page": record.get("page"),
                        "section": record.get("section") or "Document",
                        "text": chunk_text,
                    }
                )

            if end >= len(text):
                break

            next_start = end - overlap
            start = next_start if next_start > start else end

    return chunks
